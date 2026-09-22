#!/usr/bin/env python3
"""Build a 3-page quiz document from a verifica JSON file.

Page 1: 5 multiple-choice questions
Page 2: 5 open questions with ruled answer lines
Page 3: answer key (correct letter + explanation, model answers)

Usage:
    python build_verifica.py verifica.json            # -> Verifica-<Title>.docx
    python build_verifica.py verifica.json --pdf      # -> also .pdf
    python build_verifica.py verifica.json -o out.docx
"""

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OPTION_LETTERS = "ABCDEFGH"

# Successors of w:pBdr inside w:pPr, so the border element is inserted in
# schema order no matter what was set on the paragraph before.
_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)

LABELS = {
    "it": {
        "mc_title": "Domande a scelta multipla",
        "open_title": "Domande aperte",
        "key_title": "Soluzioni",
        "name": "Nome", "class": "Classe", "date": "Data",
        "mc_hint": "Segna con una crocetta l'unica risposta corretta.",
        "open_hint": "Rispondi sulle righe usando frasi complete.",
        "key_mc": "Scelta multipla", "key_open": "Domande aperte",
        "key_note": "Foglio per l'insegnante - non consegnare all'alunno.",
        "score": "Punteggio:", "points": "punti",
    },
    "en": {
        "mc_title": "Multiple choice questions",
        "open_title": "Open questions",
        "key_title": "Answer key",
        "name": "Name", "class": "Class", "date": "Date",
        "mc_hint": "Tick the one correct answer.",
        "open_hint": "Answer on the lines using full sentences.",
        "key_mc": "Multiple choice", "key_open": "Open questions",
        "key_note": "Teacher's sheet - do not hand out to the pupil.",
        "score": "Score:", "points": "points",
    },
}


def labels_for(lang):
    return LABELS.get((lang or "en").lower()[:2], LABELS["en"])


# --------------------------------------------------------------------------
# input validation
# --------------------------------------------------------------------------

def load_quiz(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    errors = []

    if not data.get("title"):
        errors.append("missing 'title'")

    mc = data.get("multiple_choice") or []
    op = data.get("open_questions") or []
    if len(mc) != 5:
        errors.append("'multiple_choice' has %d items, expected 5" % len(mc))
    if len(op) != 5:
        errors.append("'open_questions' has %d items, expected 5" % len(op))

    for i, q in enumerate(mc, 1):
        opts = q.get("options") or []
        if len(opts) < 2:
            errors.append("multiple_choice[%d]: needs at least 2 options" % i)
        if not isinstance(q.get("answer"), int) or not 0 <= q.get("answer", -1) < len(opts):
            errors.append("multiple_choice[%d]: 'answer' must be a 0-based index into options" % i)
        if not q.get("question"):
            errors.append("multiple_choice[%d]: missing 'question'" % i)
        if not q.get("explanation"):
            errors.append("multiple_choice[%d]: missing 'explanation'" % i)

    for i, q in enumerate(op, 1):
        if not q.get("question"):
            errors.append("open_questions[%d]: missing 'question'" % i)
        if not q.get("model_answer"):
            errors.append("open_questions[%d]: missing 'model_answer'" % i)

    if errors:
        sys.exit("Invalid quiz file:\n  - " + "\n  - ".join(errors))

    if len(set(q["answer"] for q in mc)) == 1:
        print("WARNING: every correct answer is the same letter - shuffle the options.",
              file=sys.stderr)

    return data


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _ruled(paragraph):
    """Give a paragraph a thin bottom border, so it reads as a writing line."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    pBdr.append(bottom)
    pPr.insert_element_before(pBdr, *_PBDR_SUCCESSORS)
    return paragraph


def _para(doc, text="", size=11, bold=False, italic=False, before=0, after=4,
          color=None, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def _page_header(doc, lab, title, subtitle, page_title, with_fields):
    _para(doc, title.upper(), size=17, bold=True, after=0)
    if subtitle:
        _para(doc, subtitle, size=10, italic=True, color="555555", after=6)
    if with_fields:
        line = "%s: ______________________   %s: __________   %s: ____________" % (
            lab["name"], lab["class"], lab["date"])
        _para(doc, line, size=10, color="555555", after=10)
    _para(doc, page_title, size=13, bold=True, before=2, after=2)


def build_docx(data, out_path):
    lab = labels_for(data.get("language"))
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    title = data["title"]
    subtitle = data.get("subject", "")

    # --- page 1: multiple choice ------------------------------------------
    _page_header(doc, lab, title, subtitle, lab["mc_title"], with_fields=True)
    _para(doc, lab["mc_hint"], size=10, italic=True, color="555555", after=8)

    for i, q in enumerate(data["multiple_choice"], 1):
        _para(doc, "%d. %s" % (i, q["question"]), size=11, bold=True, before=6, after=2)
        for j, opt in enumerate(q["options"]):
            _para(doc, "%s) %s" % (OPTION_LETTERS[j], opt), size=11, indent=0.8, after=1)

    # --- page 2: open questions -------------------------------------------
    doc.add_page_break()
    _page_header(doc, lab, title, subtitle, lab["open_title"], with_fields=True)
    _para(doc, lab["open_hint"], size=10, italic=True, color="555555", after=8)

    for i, q in enumerate(data["open_questions"], 1):
        _para(doc, "%d. %s" % (i, q["question"]), size=11, bold=True, before=6, after=4)
        for _ in range(int(q.get("lines", 3))):
            _ruled(_para(doc, "", size=11, after=8, indent=0.3))

    # --- page 3: answer key -----------------------------------------------
    doc.add_page_break()
    _page_header(doc, lab, title, subtitle, lab["key_title"], with_fields=False)
    _para(doc, lab["key_note"], size=10, italic=True, color="555555", after=8)

    _para(doc, lab["key_mc"], size=12, bold=True, before=4, after=3)
    for i, q in enumerate(data["multiple_choice"], 1):
        letter = OPTION_LETTERS[q["answer"]]
        p = _para(doc, "", size=11, after=3, indent=0.3)
        head = p.add_run("%d. %s) %s" % (i, letter, q["options"][q["answer"]]))
        head.bold = True
        head.font.size = Pt(11)
        tail = p.add_run("  -  " + q["explanation"])
        tail.font.size = Pt(10)
        tail.font.color.rgb = RGBColor.from_string("444444")

    _para(doc, lab["key_open"], size=12, bold=True, before=10, after=3)
    for i, q in enumerate(data["open_questions"], 1):
        _para(doc, "%d. %s" % (i, q["question"]), size=10, bold=True, after=1, indent=0.3)
        _para(doc, q["model_answer"], size=10, color="444444", after=5, indent=0.6)

    total = len(data["multiple_choice"]) + len(data["open_questions"])
    _para(doc, "%s ____ / %d %s" % (lab["score"], total, lab["points"]),
          size=10, italic=True, color="555555", before=10)

    doc.save(str(out_path))
    return out_path


# --------------------------------------------------------------------------
# PDF (optional)
# --------------------------------------------------------------------------

def build_pdf(data, out_path):
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer)

    lab = labels_for(data.get("language"))
    base = getSampleStyleSheet()["Normal"]
    S = {
        "title": ParagraphStyle("t", base, fontName="Helvetica-Bold", fontSize=16, leading=19, spaceAfter=2),
        "sub": ParagraphStyle("s", base, fontName="Helvetica-Oblique", fontSize=9, leading=12, textColor="#555555", spaceAfter=6),
        "fields": ParagraphStyle("f", base, fontSize=9.5, leading=13, textColor="#555555", spaceAfter=10),
        "h2": ParagraphStyle("h", base, fontName="Helvetica-Bold", fontSize=12.5, leading=15, spaceBefore=6, spaceAfter=3),
        "hint": ParagraphStyle("hi", base, fontName="Helvetica-Oblique", fontSize=9.5, leading=12, textColor="#555555", spaceAfter=8),
        "q": ParagraphStyle("q", base, fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=6, spaceAfter=2),
        "opt": ParagraphStyle("o", base, fontSize=10.5, leading=14, leftIndent=0.8 * cm, spaceAfter=1),
        "ans": ParagraphStyle("a", base, fontSize=10, leading=13, leftIndent=0.6 * cm, textColor="#444444", spaceAfter=5),
        "key": ParagraphStyle("k", base, fontSize=10.5, leading=14, leftIndent=0.3 * cm, spaceAfter=3),
    }

    def header(story, page_title, with_fields):
        story.append(Paragraph(escape(data["title"].upper()), S["title"]))
        if data.get("subject"):
            story.append(Paragraph(escape(data["subject"]), S["sub"]))
        if with_fields:
            story.append(Paragraph(
                "%s: ______________________&nbsp;&nbsp;&nbsp;"
                "%s: __________&nbsp;&nbsp;&nbsp;"
                "%s: ____________" % (lab["name"], lab["class"], lab["date"]),
                S["fields"]))
        story.append(Paragraph(escape(page_title), S["h2"]))

    story = []
    header(story, lab["mc_title"], True)
    story.append(Paragraph(escape(lab["mc_hint"]), S["hint"]))
    for i, q in enumerate(data["multiple_choice"], 1):
        story.append(Paragraph("%d. %s" % (i, escape(q["question"])), S["q"]))
        for j, opt in enumerate(q["options"]):
            story.append(Paragraph("%s) %s" % (OPTION_LETTERS[j], escape(opt)), S["opt"]))

    story.append(PageBreak())
    header(story, lab["open_title"], True)
    story.append(Paragraph(escape(lab["open_hint"]), S["hint"]))
    for i, q in enumerate(data["open_questions"], 1):
        story.append(Paragraph("%d. %s" % (i, escape(q["question"])), S["q"]))
        for _ in range(int(q.get("lines", 3))):
            story.append(Spacer(1, 0.55 * cm))
            story.append(HRFlowable(width="100%", thickness=0.4, color="#AAAAAA"))

    story.append(PageBreak())
    header(story, lab["key_title"], False)
    story.append(Paragraph(escape(lab["key_note"]), S["hint"]))
    story.append(Paragraph(escape(lab["key_mc"]), S["h2"]))
    for i, q in enumerate(data["multiple_choice"], 1):
        letter = OPTION_LETTERS[q["answer"]]
        story.append(Paragraph(
            "<b>%d. %s) %s</b><font size=9 color='#444444'> &ndash; %s</font>" % (
                i, letter, escape(q["options"][q["answer"]]), escape(q["explanation"])),
            S["key"]))
    story.append(Paragraph(escape(lab["key_open"]), S["h2"]))
    for i, q in enumerate(data["open_questions"], 1):
        story.append(Paragraph("<b>%d. %s</b>" % (i, escape(q["question"])), S["key"]))
        story.append(Paragraph(escape(q["model_answer"]), S["ans"]))

    total = len(data["multiple_choice"]) + len(data["open_questions"])
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph(
        "%s ____ / %d %s" % (lab["score"], total, lab["points"]), S["hint"]))

    SimpleDocTemplate(
        str(out_path), pagesize=A4,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        leftMargin=2.0 * cm, rightMargin=2.0 * cm,
        title=data["title"], author="appunti-lezione",
    ).build(story)
    return out_path


# --------------------------------------------------------------------------

def slugify(text):
    text = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE).strip()
    return re.sub(r"[\s_-]+", "-", text) or "verifica"


def main():
    ap = argparse.ArgumentParser(
        description="Build a 3-page quiz document from a verifica JSON file.")
    ap.add_argument("quiz_json", help="path to the verifica JSON file")
    ap.add_argument("-o", "--output",
                    help="output .docx path (default: Verifica-<Title>.docx next to the JSON)")
    ap.add_argument("--pdf", action="store_true",
                    help="also render a PDF with the same base name")
    args = ap.parse_args()

    src = Path(args.quiz_json)
    data = load_quiz(src)

    if args.output:
        docx_path = Path(args.output)
    else:
        docx_path = src.parent / ("Verifica-%s.docx" % slugify(data["title"]))
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    build_docx(data, docx_path)
    print("DOCX: %s" % docx_path)

    if args.pdf:
        pdf_path = docx_path.with_suffix(".pdf")
        build_pdf(data, pdf_path)
        print("PDF:  %s" % pdf_path)


if __name__ == "__main__":
    main()
